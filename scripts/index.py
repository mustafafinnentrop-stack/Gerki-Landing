#!/usr/bin/env python3
"""
Gerki Index
Liest Dokumente aus konfigurierten Ordnern und erstellt einen durchsuchbaren Index.
Unterstützt: .txt, .md, .csv, .json, .xml, .eml
Für PDF/DOCX/XLSX: optionale Abhängigkeiten (Hinweis wird angezeigt)
"""

import json
import os
import sys
import hashlib
import datetime
from pathlib import Path

CONFIG_FILE = Path.home() / ".gerki" / "config.json"
INDEX_DIR = Path.home() / ".gerki" / "index"
INDEX_FILE = INDEX_DIR / "documents.json"

SUPPORTED_PLAIN = {".txt", ".md", ".csv", ".json", ".xml", ".eml", ".msg"}
SUPPORTED_RICH = {".pdf", ".docx", ".xlsx", ".pptx"}
CHUNK_SIZE = 1500  # Zeichen pro Chunk


def load_config():
    if not CONFIG_FILE.exists():
        print("Fehler: Gerki ist noch nicht eingerichtet.")
        print("Bitte zuerst ausführen: python scripts/setup.py")
        sys.exit(1)
    with open(CONFIG_FILE, encoding="utf-8") as f:
        return json.load(f)


def load_index():
    if INDEX_FILE.exists():
        with open(INDEX_FILE, encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_index(index):
    INDEX_DIR.mkdir(parents=True, exist_ok=True)
    with open(INDEX_FILE, "w", encoding="utf-8") as f:
        json.dump(index, f, indent=2, ensure_ascii=False)


def file_hash(path):
    h = hashlib.md5()
    with open(path, "rb") as f:
        h.update(f.read())
    return h.hexdigest()


def read_plain(path):
    try:
        with open(path, encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception:
        return None


def read_pdf(path):
    try:
        import pdfplumber
        text = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text.append(t)
        return "\n".join(text)
    except ImportError:
        return f"[PDF nicht lesbar – installiere: pip install pdfplumber]"
    except Exception as e:
        return f"[Fehler beim Lesen: {e}]"


def read_docx(path):
    try:
        import docx
        doc = docx.Document(path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except ImportError:
        return "[DOCX nicht lesbar – installiere: pip install python-docx]"
    except Exception as e:
        return f"[Fehler beim Lesen: {e}]"


def read_xlsx(path):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows = []
        for sheet in wb.worksheets:
            rows.append(f"[Tabelle: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                line = " | ".join(str(c) for c in row if c is not None)
                if line.strip():
                    rows.append(line)
        return "\n".join(rows)
    except ImportError:
        return "[XLSX nicht lesbar – installiere: pip install openpyxl]"
    except Exception as e:
        return f"[Fehler beim Lesen: {e}]"


def read_file(path):
    ext = path.suffix.lower()
    if ext in SUPPORTED_PLAIN:
        return read_plain(path)
    elif ext == ".pdf":
        return read_pdf(path)
    elif ext == ".docx":
        return read_docx(path)
    elif ext in {".xlsx", ".xls"}:
        return read_xlsx(path)
    return None


def chunk_text(text, size=CHUNK_SIZE):
    chunks = []
    for i in range(0, len(text), size):
        chunk = text[i:i + size].strip()
        if chunk:
            chunks.append(chunk)
    return chunks


def index_folder(folder, existing_index, rebuild=False):
    folder = Path(folder)
    if not folder.exists():
        print(f"  Übersprungen (nicht gefunden): {folder}")
        return {}

    new_entries = {}
    all_exts = SUPPORTED_PLAIN | SUPPORTED_RICH

    files = [f for f in folder.rglob("*") if f.is_file() and f.suffix.lower() in all_exts]

    for file in files:
        try:
            fhash = file_hash(file)
            str_path = str(file)

            # Skip wenn unverändert und nicht rebuild
            if not rebuild and str_path in existing_index:
                if existing_index[str_path].get("hash") == fhash:
                    new_entries[str_path] = existing_index[str_path]
                    continue

            text = read_file(file)
            if not text:
                continue

            chunks = chunk_text(text)
            mod_time = datetime.datetime.fromtimestamp(file.stat().st_mtime).isoformat()

            new_entries[str_path] = {
                "path": str_path,
                "name": file.name,
                "extension": file.suffix.lower(),
                "hash": fhash,
                "modified": mod_time,
                "chunks": chunks,
                "chunk_count": len(chunks),
                "char_count": len(text)
            }
            print(f"  ✓ {file.name} ({len(chunks)} Abschnitte)")

        except Exception as e:
            print(f"  ! Fehler bei {file.name}: {e}")

    return new_entries


def update_config_timestamp(config):
    config["last_indexed"] = datetime.datetime.now().isoformat()
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump(config, f, indent=2, ensure_ascii=False)


def main():
    rebuild = "--rebuild" in sys.argv
    called_directly = __name__ == "__main__"

    config = load_config()
    existing_index = load_index() if not rebuild else {}

    folders = config.get("folders", [])
    if not folders:
        print("Keine Ordner konfiguriert. Bitte setup.py ausführen.")
        if called_directly:
            sys.exit(1)
        return

    print()
    print("Gerki indexiert deine Dokumente...")
    if rebuild:
        print("Modus: Vollständiger Neuaufbau")
    print()

    full_index = {}
    total_files = 0

    for folder in folders:
        print(f"Ordner: {folder}")
        entries = index_folder(folder, existing_index, rebuild=rebuild)
        full_index.update(entries)
        total_files += len(entries)
        print()

    save_index(full_index)
    update_config_timestamp(config)

    print("=" * 40)
    print(f"  {total_files} Dateien indexiert")
    print(f"  Index gespeichert: {INDEX_FILE}")
    print("=" * 40)


if __name__ == "__main__":
    main()
